from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

NAVY = colors.HexColor("#1A2E40")
TERRACOTA = colors.HexColor("#E07A5F")
GRIS = colors.HexColor("#6B7280")
VERDE = colors.HexColor("#15803D")


def nombre_archivo_reporte_cuotas(gasto, extension):
    periodo = gasto.periodo.replace(" ", "_")
    return f"gastos_comunes_{periodo}.{extension}"


def _totales(cuotas):
    total_pagado = sum(c.monto for c in cuotas if c.estado == "pagado")
    total_pendiente = sum(c.monto for c in cuotas if c.estado != "pagado")
    return total_pagado, total_pendiente


def generar_cuotas_excel(gasto, cuotas, condominio):
    """Reporte de cuotas de un periodo, para que administración lo siga
    trabajando/filtrando en una planilla -- montos como número real (no
    texto), para que se puedan sumar/filtrar en Excel sin reescribirlos."""
    wb = Workbook()
    ws = wb.active
    ws.title = gasto.periodo[:31]  # Excel no admite nombres de hoja de más de 31 caracteres

    ws.append([f"{condominio.nombre} -- Gastos comunes, periodo {gasto.periodo}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
    ws.append([])

    encabezados = [condominio.etiqueta_torre, condominio.etiqueta_unidad, "Monto", "Estado", "Fecha de pago"]
    ws.append(encabezados)
    fila_encabezado = ws.max_row
    for col in range(1, len(encabezados) + 1):
        celda = ws.cell(row=fila_encabezado, column=col)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1A2E40")

    for c in cuotas:
        ws.append([
            c.unidad.torre.nombre if c.unidad.torre_id else "—",
            c.unidad.numero,
            c.monto,
            c.get_estado_display(),
            c.fecha_pago.strftime("%d-%m-%Y") if c.fecha_pago else "—",
        ])
        ws.cell(row=ws.max_row, column=3).number_format = "#,##0"

    total_pagado, total_pendiente = _totales(cuotas)
    ws.append([])
    for etiqueta, monto in (("Total recaudado", total_pagado), ("Total pendiente", total_pendiente)):
        ws.append(["", "", monto, etiqueta])
        ws.cell(row=ws.max_row, column=3).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=3).number_format = "#,##0"
        ws.cell(row=ws.max_row, column=4).font = Font(bold=True)

    for i, ancho in enumerate((12, 14, 14, 14, 16), start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    response = HttpResponse(buffer.read(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{nombre_archivo_reporte_cuotas(gasto, "xlsx")}"'
    return response


def generar_cuotas_pdf(gasto, cuotas, condominio):
    """Mismo reporte que generar_cuotas_excel, pero para imprimir/adjuntar en
    un correo o acta -- montos formateados como texto ($X.XXX), no pensado
    para seguir trabajándolo (para eso está el Excel)."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=A4, topMargin=1.5 * cm, bottomMargin=1.5 * cm, leftMargin=1.5 * cm, rightMargin=1.5 * cm)
    estilos = getSampleStyleSheet()
    elementos = [
        Paragraph(f"<b>{condominio.nombre}</b> -- Gastos comunes, periodo {gasto.periodo}", estilos["Title"]),
        Spacer(1, 12),
    ]

    encabezados = [condominio.etiqueta_torre, condominio.etiqueta_unidad, "Monto", "Estado", "Fecha de pago"]
    filas = [encabezados]
    for c in cuotas:
        filas.append([
            c.unidad.torre.nombre if c.unidad.torre_id else "—",
            c.unidad.numero,
            f"${c.monto:,}".replace(",", "."),
            c.get_estado_display(),
            c.fecha_pago.strftime("%d-%m-%Y") if c.fecha_pago else "—",
        ])

    tabla = Table(filas, repeatRows=1)
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, GRIS),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F9FA")]),
        ("ALIGN", (2, 0), (2, -1), "RIGHT"),
    ]))
    elementos.append(tabla)
    elementos.append(Spacer(1, 16))

    total_pagado, total_pendiente = _totales(cuotas)
    resumen = Table([
        ["Total recaudado", f"${total_pagado:,}".replace(",", ".")],
        ["Total pendiente", f"${total_pendiente:,}".replace(",", ".")],
    ])
    resumen.setStyle(TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 11),
        ("TEXTCOLOR", (1, 0), (1, 0), VERDE),
        ("TEXTCOLOR", (1, 1), (1, 1), TERRACOTA),
    ]))
    elementos.append(resumen)

    doc.build(elementos)
    buffer.seek(0)
    response = HttpResponse(buffer.read(), content_type="application/pdf")
    response["Content-Disposition"] = f'attachment; filename="{nombre_archivo_reporte_cuotas(gasto, "pdf")}"'
    return response
